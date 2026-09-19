#!/usr/bin/env python3
"""Create a source-preserving Word transcription of every EPUB spine document.

Rendered equation/code images are embedded, not guessed or silently OCR-ed.
Original XHTML/CSS and commented MathML remain in the source bundle. Word layout
is a reflowed transcription, not a facsimile of publisher pagination. The report
checks the complete visible-text sequence and every image placement independently
against the input DOM. Semantic/mathematical review is a separate gate.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any
from zipfile import ZipFile
from urllib.parse import urlsplit, unquote
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image

from .inventory_epub import EPUB, NS, HEADING_TAGS, local_name, resolve_member, visible_text, write_json

BLOCK_TAGS = {"p", "pre", "table", "ol", "ul", "figure", "section", "div", "header", "nav", "aside", "blockquote", "figcaption"} | HEADING_TAGS


class WordExporter:
    def __init__(self, root: Path) -> None:
        self.root = root
        self.source = root / "instructions" / "01-source"
        self.unpacked = self.source / "private" / "epub"
        self.document = Document()
        self.document.core_properties.title = "Advances in Financial Machine Learning — source transcription"
        self.document.core_properties.author = "Marcos López de Prado (source); transcription for Oscar"
        self.document.core_properties.comments = "Reflowed EPUB transcription. Source images retained. See extraction audit; no claim of mathematical verification."
        section = self.document.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.75)
        section.header_distance = section.footer_distance = Inches(0.3)
        normal = self.document.styles["Normal"]
        normal.font.name, normal.font.size = "Cambria", Pt(10.5)
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.line_spacing = 1.05
        for level in range(1, 7):
            style = self.document.styles[f"Heading {level}"]
            style.font.name = "Cambria"
            style.font.size = Pt({1: 21, 2: 15, 3: 12, 4: 11, 5: 11, 6: 10.5}[level])
            style.font.color.rgb = RGBColor.from_string("17365D")
            style.paragraph_format.space_before = Pt(12 if level > 1 else 0)
            style.paragraph_format.space_after = Pt(7)
            style.paragraph_format.keep_with_next = True
        for name in ["List Bullet", "List Number"]:
            self.document.styles[name].font.name = "Cambria"
            self.document.styles[name].font.size = Pt(10.5)
        header = section.header.paragraphs[0]
        header.text = "ADVANCES IN FINANCIAL MACHINE LEARNING  |  SOURCE TRANSCRIPTION"
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string("666666")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Word page ").font.size = Pt(8)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        self.image_events: list[dict[str, Any]] = []
        self.page_bookmarks: list[dict[str, str]] = []
        self.current_member = ""
        self.current_index_document = False
        self.bookmark_number = 0
        self.numbering_number = 100
        self.anchor_names: set[str] = set()
        self.hyperlink_events: list[dict[str, str]] = []

    def source_anchor(self, paragraph: Any, anchor: str) -> str:
        name = "src_" + hashlib.sha256((self.current_member + "#" + anchor).encode()).hexdigest()[:24]
        if name in self.anchor_names:
            return name
        self.anchor_names.add(name)
        self.bookmark_number += 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self.bookmark_number))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self.bookmark_number))
        paragraph._p.append(start)
        paragraph._p.append(end)
        return name

    def bookmark(self, paragraph: Any, label: str, anchor: str) -> None:
        name = self.source_anchor(paragraph, anchor)
        self.page_bookmarks.append({"source_page": label, "bookmark": name})

    def list_numbering(self, node: etree._Element, depth: int) -> int | None:
        classes = node.get("class", "").split()
        ancestors = list(node.iterancestors())
        is_navigation = any(local_name(parent) == "nav" for parent in ancestors)
        is_index = self.current_index_document
        if is_navigation or is_index or {"none", "biblioEntryList"}.intersection(classes):
            return None
        self.numbering_number += 1
        number_id = self.numbering_number
        format_name = "decimal" if local_name(node) == "ol" else "bullet"
        if local_name(node) == "ol":
            ordered_ancestors = sum(local_name(parent) == "ol" for parent in ancestors)
            if ordered_ancestors:
                format_name = "lowerLetter" if ordered_ancestors == 1 else "lowerRoman"
        for source_class, word_format in [("lower-latin", "lowerLetter"), ("upper-latin", "upperLetter"),
                                           ("lower-roman", "lowerRoman"), ("upper-roman", "upperRoman")]:
            if source_class in classes:
                format_name = word_format
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(number_id))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in [("start", node.get("start", "1")), ("numFmt", format_name),
                           ("lvlText", ("○" if "circle" in classes else "•") if format_name == "bullet" else "%1.")]:
            element = OxmlElement("w:" + tag)
            element.set(qn("w:val"), value)
            level.append(element)
        abstract.append(level)
        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(number_id))
        identifier = OxmlElement("w:abstractNumId")
        identifier.set(qn("w:val"), str(number_id))
        number.append(identifier)
        numbering = self.document.part.numbering_part.element
        numbering.append(abstract)
        numbering.append(number)
        return number_id

    def add_image(self, paragraph: Any, node: etree._Element, *, inline: bool) -> None:
        asset = resolve_member(self.current_member, node.get("src", ""))
        if asset is None:
            raise ValueError("External source images are not permitted in this export.")
        raw = (self.unpacked / asset).read_bytes()
        with Image.open(io.BytesIO(raw)) as image:
            width_px, height_px = image.size
            format_name = image.format
            pixels_sha256 = hashlib.sha256(image.convert("RGBA").tobytes()).hexdigest()
            if format_name == "GIF":
                converted = io.BytesIO()
                image.convert("RGBA").save(converted, format="PNG")
                image_data = converted.getvalue()
            else:
                image_data = raw
        basename = Path(asset).name
        is_equation = "_M" in basename or "_ILM" in basename
        if inline or "_ILM" in basename:
            height = min(height_px / 144.0, 0.38)
            width = height * width_px / height_px
        elif is_equation:
            width = min(width_px / 144.0, 6.8)
            height = width * height_px / width_px
        else:
            width = min(width_px / 120.0, 6.8)
            height = width * height_px / width_px
        # Never crop or distort an equation, figure, or code image.
        factor = min(1.0, 6.8 / width, 8.7 / height)
        width, height = width * factor, height * factor
        run = paragraph.add_run()
        picture = run.add_picture(io.BytesIO(image_data), width=Inches(width), height=Inches(height))
        picture._inline.docPr.set("descr", asset)
        picture._inline.docPr.set("title", node.get("alt", asset))
        if not inline:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
        self.image_events.append({"asset": asset, "source_bytes_sha256": hashlib.sha256(raw).hexdigest(),
                                  "embedded_bytes_sha256": hashlib.sha256(image_data).hexdigest(),
                                  "pixels_sha256": pixels_sha256, "width_inches": width, "height_inches": height,
                                  "lossless_format_conversion": format_name == "GIF"})

    @staticmethod
    def add_text(paragraph: Any, text: str | None, flags: dict[str, bool]) -> None:
        if not text:
            return
        run = paragraph.add_run(text)
        run.bold = flags.get("bold", False)
        run.italic = flags.get("italic", False)
        run.font.subscript = flags.get("subscript", False)
        run.font.superscript = flags.get("superscript", False)
        if flags.get("code"):
            run.font.name = "Consolas"
            run.font.size = Pt(9)

    def inline(self, node: etree._Element, paragraph: Any, flags: dict[str, bool] | None = None) -> None:
        if not isinstance(node.tag, str):
            return
        flags = dict(flags or {})
        tag = local_name(node)
        for candidate, flag in [("b", "bold"), ("strong", "bold"), ("i", "italic"), ("em", "italic"),
                                ("sub", "subscript"), ("sup", "superscript"), ("code", "code")]:
            if tag == candidate:
                flags[flag] = True
        if "codeLabel" in node.get("class", "").split():
            flags["code"] = True
        if node.get(f"{{{EPUB}}}type") == "pagebreak":
            self.bookmark(paragraph, node.get("title") or node.get("id", "unknown"), node.get("id", "unknown"))
        if node.get("id"):
            self.source_anchor(paragraph, node.get("id"))
        hyperlink = None
        first_child_index = len(paragraph._p)
        if tag == "a" and node.get("href"):
            href = node.get("href")
            link = urlsplit(href)
            hyperlink = OxmlElement("w:hyperlink")
            if link.scheme or link.netloc:
                relationship = self.document.part.relate_to(href, RT.HYPERLINK, is_external=True)
                hyperlink.set(qn("r:id"), relationship)
            else:
                target = resolve_member(self.current_member, href)
                anchor = unquote(link.fragment)
                name = "src_" + hashlib.sha256((str(target) + "#" + anchor).encode()).hexdigest()[:24]
                hyperlink.set(qn("w:anchor"), name)
            self.hyperlink_events.append({"source_document": self.current_member, "href": href})
        if tag == "img":
            self.add_image(paragraph, node, inline=True)
            return
        if tag == "br":
            paragraph.add_run().add_break()
            return
        self.add_text(paragraph, node.text, flags)
        for child in node:
            self.inline(child, paragraph, flags)
            self.add_text(paragraph, child.tail, flags)
        if hyperlink is not None:
            for child in list(paragraph._p)[first_child_index:]:
                paragraph._p.remove(child)
                hyperlink.append(child)
            paragraph._p.append(hyperlink)

    def list_item(self, node: etree._Element, container: Any, number_id: int | None, depth: int) -> None:
        paragraph = None
        first_content = True

        def new_paragraph() -> Any:
            nonlocal first_content
            result = container.add_paragraph()
            left = depth * 0.20 + (0.24 if number_id is not None else 0.0)
            result.paragraph_format.left_indent = Inches(left)
            result.paragraph_format.first_line_indent = Inches(-0.24 if number_id is not None and first_content else 0.0)
            if number_id is not None:
                result.paragraph_format.tab_stops.add_tab_stop(Inches(left))
            if first_content and number_id is not None:
                properties = result._p.get_or_add_pPr()
                numbering = OxmlElement("w:numPr")
                for tag, value in [("ilvl", "0"), ("numId", str(number_id))]:
                    element = OxmlElement("w:" + tag)
                    element.set(qn("w:val"), value)
                    numbering.append(element)
                properties.append(numbering)
            first_content = False
            return result

        if node.text and node.text.strip():
            paragraph = new_paragraph()
            self.add_text(paragraph, node.text, {})
        for child in node:
            tag = local_name(child)
            if tag in {"ol", "ul"}:
                self.block(child, container, depth=depth + 1)
                paragraph = None
            elif tag == "p":
                paragraph = new_paragraph()
                self.inline(child, paragraph)
            elif tag in {"div", "figure", "table", "pre"} | HEADING_TAGS:
                self.block(child, container, depth=depth)
                paragraph = None
            elif tag != "#comment":
                if paragraph is None:
                    paragraph = new_paragraph()
                self.inline(child, paragraph)
            if child.tail and child.tail.strip():
                if paragraph is None:
                    paragraph = new_paragraph()
                self.add_text(paragraph, child.tail, {})

    def table(self, node: etree._Element, container: Any) -> None:
        rows = node.findall(".//h:tr", NS)
        column_count = max(sum(int(cell.get("colspan", "1")) for cell in row if local_name(cell) in {"td", "th"}) for row in rows)
        table = container.add_table(rows=len(rows), cols=column_count)
        table.style = "Table Grid"
        table.autofit = True
        for row_index, row in enumerate(rows):
            column = 0
            for cell_node in row:
                if local_name(cell_node) not in {"td", "th"}:
                    continue
                span = int(cell_node.get("colspan", "1"))
                cell = table.cell(row_index, column)
                if span > 1:
                    cell = cell.merge(table.cell(row_index, column + span - 1))
                # Use the existing empty paragraph for inline-only cell content.
                if any(local_name(child) in BLOCK_TAGS for child in cell_node):
                    self.block(cell_node, cell)
                    if not cell.paragraphs[0].text:
                        cell._tc.remove(cell.paragraphs[0]._p)
                else:
                    self.inline(cell_node, cell.paragraphs[0])
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5 if column_count < 9 else 7.5)
                column += span
            if row_index == 0:
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)

    def block(self, node: etree._Element, container: Any, depth: int = 0) -> None:
        if not isinstance(node.tag, str):
            return
        tag = local_name(node)
        if tag in HEADING_TAGS:
            paragraph = container.add_paragraph(style=f"Heading {int(tag[1:])}")
            self.inline(node, paragraph)
            return
        if tag in {"p", "pre"}:
            paragraph = container.add_paragraph()
            if tag == "pre":
                paragraph.paragraph_format.space_after = Pt(4)
            self.inline(node, paragraph, {"code": tag == "pre"})
            return
        if tag == "img":
            self.add_image(container.add_paragraph(), node, inline=False)
            return
        if tag == "table":
            self.table(node, container)
            return
        if tag in {"ul", "ol"}:
            number_id = self.list_numbering(node, depth)
            for child in node:
                if local_name(child) == "li":
                    self.list_item(child, container, number_id, depth)
                else:
                    self.block(child, container, depth)
            return
        # Container anchors are represented by empty bookmarks, not new text.
        if node.get("id") or tag == "body":
            self.source_anchor(container.add_paragraph(), node.get("id", ""))
        # Containers carry no implicit replacement text. All non-whitespace text
        # and child tails are passed through once in the original reading order.
        paragraph = None
        if node.text and node.text.strip():
            paragraph = container.add_paragraph()
            self.add_text(paragraph, node.text, {})
        if node.get(f"{{{EPUB}}}type") == "pagebreak":
            if paragraph is None:
                paragraph = container.add_paragraph()
            self.bookmark(paragraph, node.get("title") or node.get("id", "unknown"), node.get("id", "unknown"))
        for child in node:
            child_tag = local_name(child)
            if child_tag in BLOCK_TAGS or child_tag in {"img", "body", "td"}:
                self.block(child, container, depth)
                paragraph = None
            elif child_tag != "#comment":
                if paragraph is None:
                    paragraph = container.add_paragraph()
                self.inline(child, paragraph)
            if child.tail and child.tail.strip():
                if paragraph is None:
                    paragraph = container.add_paragraph()
                self.add_text(paragraph, child.tail, {})

    def export(self, output: Path) -> dict[str, Any]:
        spine = json.loads((self.source / "spine.json").read_text())
        expected_text: list[str] = []
        expected_images: list[str] = []
        for index, item in enumerate(spine):
            self.current_member = item["path"]
            tree = etree.parse(str(self.unpacked / self.current_member), etree.XMLParser(resolve_entities=False, no_network=True))
            self.current_index_document = any(Path(link.get("href", "")).name == "index.css" for link in tree.findall("h:head/h:link", NS))
            body = tree.find("h:body", NS)
            if body is None:
                raise ValueError("Missing source body.")
            if index:
                self.document.add_page_break()
            expected_text.append(visible_text(body))
            expected_images.extend(resolve_member(self.current_member, image.get("src", "")) for image in body.findall(".//h:img", NS))
            self.block(body, self.document)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(output)
        with ZipFile(output) as archive:
            word = etree.fromstring(archive.read("word/document.xml"))
            actual_text = "".join(word.xpath('//*[local-name()="t" and namespace-uri()="http://schemas.openxmlformats.org/wordprocessingml/2006/main"]/text()'))
            actual_images = word.xpath('//*[local-name()="docPr"]/@descr')
            embeds = word.xpath('//*[local-name()="blip"]/@*[local-name()="embed"]')
            rel_tree = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
            targets = {rel.get("Id"): rel.get("Target") for rel in rel_tree}
            image_hashes = [hashlib.sha256(archive.read("word/" + targets[identifier])).hexdigest() for identifier in embeds]
        compact = lambda value: re.sub(r"\s+", "", value)
        source_text = "".join(expected_text)
        report = {"file": output.name, "docx_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
                  "source_characters_excluding_whitespace": len(compact(source_text)),
                  "word_characters_excluding_whitespace": len(compact(actual_text)),
                  "entire_visible_text_sequence_equal_excluding_whitespace": compact(source_text) == compact(actual_text),
                  "source_image_placements": len(expected_images), "word_image_placements": len(actual_images),
                  "entire_image_placement_sequence_equal": expected_images == actual_images,
                  "all_embedded_image_bytes_match_export_manifest": image_hashes == [event["embedded_bytes_sha256"] for event in self.image_events],
                  "page_bookmark_count": len(self.page_bookmarks),
                  "hyperlink_count": len(self.hyperlink_events),
                  "source_anchor_count": len(self.anchor_names),
                  "mathematical_verification": "pending_independent_review",
                  "editable_equations": False,
                  "publisher_layout_fidelity": "Original EPUB retained; Word is reflowed, not a facsimile.",
                  "visual_review": "pending_render_and_review"}
        write_json(self.source / "word_export_report.json", report)
        write_json(self.source / "word_image_manifest.json", self.image_events)
        write_json(self.source / "word_page_bookmarks.json", self.page_bookmarks)
        write_json(self.source / "word_hyperlinks.json", self.hyperlink_events)
        if not report["entire_visible_text_sequence_equal_excluding_whitespace"]:
            (self.source / "private" / "word_actual_text.txt").write_text(actual_text)
            raise ValueError("Word text-sequence validation failed; inspect report before using output.")
        if not report["entire_image_placement_sequence_equal"]:
            raise ValueError("Word image-placement validation failed.")
        print(json.dumps(report, indent=2))
        return report

# CLI entry point: instructions/tools/rebuild_private_source.py.
